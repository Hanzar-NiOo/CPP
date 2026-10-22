#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Shared PowerPoint design system / toolkit for the 42 C++ module decks.

Contains the whole visual language: colours, typography, the C++ syntax
highlighter, bullets, callouts, tables, UML class boxes and inheritance arrows.

Requires: python-pptx   (Pillow is optional, used for accurate text measuring)
"""

import os
import re

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.enum.shapes import MSO_SHAPE, MSO_CONNECTOR
from pptx.enum.dml import MSO_LINE_DASH_STYLE
from pptx.oxml import parse_xml
from pptx.oxml.ns import nsdecls, qn

# --------------------------------------------------------------------------
# Design system
# --------------------------------------------------------------------------
NAVY     = RGBColor(0x0E, 0x1B, 0x2E)
NAVY_2   = RGBColor(0x18, 0x2B, 0x47)
NAVY_3   = RGBColor(0x25, 0x3D, 0x60)
ORANGE   = RGBColor(0xE8, 0x87, 0x1A)
ORANGE_L = RGBColor(0xFD, 0xF0, 0xDF)
TEAL     = RGBColor(0x1F, 0x8A, 0x7A)
TEAL_L   = RGBColor(0xE3, 0xF3, 0xF1)
RED      = RGBColor(0xC0, 0x39, 0x2B)
RED_L    = RGBColor(0xFB, 0xEA, 0xE8)
WHITE    = RGBColor(0xFF, 0xFF, 0xFF)
PAGE     = RGBColor(0xF6, 0xF8, 0xFB)
INK      = RGBColor(0x1C, 0x27, 0x38)
GREY     = RGBColor(0x5E, 0x6E, 0x84)
GREY_L   = RGBColor(0x9A, 0xA8, 0xBA)
LINE     = RGBColor(0xD9, 0xDF, 0xEA)
CODE_BG  = RGBColor(0x14, 0x1B, 0x27)
TERM_BG  = RGBColor(0x0B, 0x12, 0x0B)

C_CODE = {
    "def":   RGBColor(0xD6, 0xDE, 0xEA),
    "kw":    RGBColor(0xC7, 0x92, 0xEA),
    "type":  RGBColor(0x82, 0xAA, 0xFF),
    "str":   RGBColor(0xEC, 0xC4, 0x8D),
    "cmt":   RGBColor(0x62, 0x73, 0x8A),
    "pre":   RGBColor(0xF0, 0x71, 0x78),
    "num":   RGBColor(0xF7, 0x8C, 0x6C),
    "fn":    RGBColor(0xFF, 0xCB, 0x6B),
}

KEYWORDS = {
    "class", "public", "private", "protected", "virtual", "return", "void",
    "const", "if", "else", "for", "while", "new", "delete", "operator",
    "this", "using", "namespace", "static", "bool", "true", "false", "unsigned",
    "int", "std", "string", "include", "define", "ifndef", "endif", "struct",
    "explicit", "override", "final", "nullptr", "mutable", "typename",
}

FONT   = "Segoe UI"
FONT_B = "Segoe UI Semibold"
MONO   = "Consolas"

SW = Inches(13.333)
SH = Inches(7.5)
ML = Inches(0.85)
MR = Inches(0.85)
CW = SW - ML - MR
BODY_TOP = Inches(1.58)
FOOT_TOP = SH - Inches(0.72)          # usable bottom of the body area

OUT_NAME = os.environ.get("PPTX_OUT")

# Caption printed in the footer of every content slide.  Each module sets
# its own value right after importing the theme.
FOOTER_LABEL = "CPP Module 03  ·  Inheritance, Interfaces & Polymorphism"

# line height factors (em) used for layout maths
LH_TEXT = 1.50      # Segoe UI, line_spacing 1.12
LH_CODE = 1.42      # Consolas, line_spacing 1.18


# --------------------------------------------------------------------------
# Text measuring (PIL when available, otherwise estimated)
# --------------------------------------------------------------------------
try:
    from PIL import ImageFont as _PILFont
    _HAVE_PIL = True
except Exception:
    _HAVE_PIL = False

_FONT_FILES = {
    "Segoe UI":         ("segoeui.ttf",  "segoeuib.ttf"),
    "Segoe UI Semibold": ("seguisb.ttf",  "seguisb.ttf"),
    "Consolas":         ("consola.ttf",  "consolab.ttf"),
}
_FONT_DIRS = ["C:/Windows/Fonts", "/mnt/c/Windows/Fonts"]
_fcache = {}


def _font_file(name, bold):
    reg, bld = _FONT_FILES.get(name, _FONT_FILES["Segoe UI"])
    fn = bld if bold else reg
    for d in _FONT_DIRS:
        p = os.path.join(d, fn)
        if os.path.exists(p):
            return p
    return None


def _pil_font(name, size_pt, bold=False):
    key = (name, round(size_pt, 1), bold)
    if key in _fcache:
        return _fcache[key]
    p = _font_file(name, bold)
    f = None
    if p:
        try:
            f = _PILFont.truetype(p, max(1, int(round(size_pt * 96 / 72.0))),
                                  layout_engine=_PILFont.Layout.BASIC)
        except Exception:
            f = None
    _fcache[key] = f
    return f


def text_px(txt, size_pt, name=FONT, bold=False):
    """Width of txt in px (96 dpi)."""
    if not txt:
        return 0.0
    f = _pil_font(name, size_pt, bold)
    if f is not None:
        try:
            return f.getlength(txt)
        except Exception:
            pass
    ratio = 0.55 if name == MONO else 0.50
    return len(txt) * size_pt * (96 / 72.0) * ratio


INLINE_RE = re.compile(r"(\*\*.+?\*\*|`[^`]+`)")


def plain_runs(txt, size, font=FONT, bold=False):
    """Markdown-ish text -> [(text, size, fontname, bold)] for measuring."""
    out = []
    for tok in INLINE_RE.split(txt):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            out.append((tok[2:-2], size, font, True))
        elif tok.startswith("`") and tok.endswith("`"):
            out.append((tok[1:-1], size * 0.94, MONO, False))
        else:
            out.append((tok, size, font, bold))
    return out


def count_lines(runs, max_px):
    """Number of wrapped lines for a list of (text, size, font, bold) runs."""
    if max_px <= 4:
        return 1
    lines = 1
    cur = 0.0
    for txt, sz, nm, bd in runs:
        toks = []
        buf = ""
        for ch in txt:
            buf += ch
            if ch == " ":
                toks.append(buf)
                buf = ""
        if buf:
            toks.append(buf)
        for tok in toks:
            w = text_px(tok, sz, nm, bd)
            if cur + w > max_px and cur > 0:
                lines += 1
                cur = 0.0
                tok = tok.lstrip(" ")
                w = text_px(tok, sz, nm, bd)
            cur += w
    return lines


# --------------------------------------------------------------------------
# Low level helpers
# --------------------------------------------------------------------------
def no_shadow(shape):
    try:
        shape.shadow.inherit = False
    except Exception:
        pass


def rect(slide, x, y, w, h, fill=None, line=None, lw=1.0,
         shape=MSO_SHAPE.RECTANGLE, adj=None):
    sp = slide.shapes.add_shape(shape, int(x), int(y), int(w), int(h))
    if fill is None:
        sp.fill.background()
    else:
        sp.fill.solid()
        sp.fill.fore_color.rgb = fill
    if line is None:
        sp.line.fill.background()
    else:
        sp.line.color.rgb = line
        sp.line.width = Pt(lw)
    if adj is not None:
        try:
            sp.adjustments[0] = adj
        except Exception:
            pass
    no_shadow(sp)
    sp.text_frame.word_wrap = True
    return sp


def tbox(slide, x, y, w, h):
    tb = slide.shapes.add_textbox(int(x), int(y), int(w), int(h))
    tf = tb.text_frame
    tf.word_wrap = True
    tf.margin_left = 0
    tf.margin_right = 0
    tf.margin_top = 0
    tf.margin_bottom = 0
    return tf


def new_p(tf):
    p0 = tf.paragraphs[0]
    if len(tf.paragraphs) == 1 and not p0.runs and p0.text == "":
        return p0
    return tf.add_paragraph()


def set_bullet(p, char, color, marL=0.26, font="Arial"):
    pPr = p._p.get_or_add_pPr()
    pPr.set("marL", str(int(marL * 914400)))
    pPr.set("indent", str(int(-marL * 914400)))
    for xml in (
        '<a:buClr %s><a:srgbClr val="%s"/></a:buClr>' % (nsdecls("a"), color),
        '<a:buFont %s typeface="%s"/>' % (nsdecls("a"), font),
        '<a:buChar %s char="%s"/>' % (nsdecls("a"), char),
    ):
        pPr.append(parse_xml(xml))


def no_bullet(p):
    p._p.get_or_add_pPr().append(parse_xml('<a:buNone %s/>' % nsdecls("a")))


def _line_break(p):
    if hasattr(p, "add_line_break"):
        p.add_line_break()
    else:
        p._p.append(parse_xml('<a:br %s/>' % nsdecls("a")))


def _emit(p, txt, size, color, font, bold, italic, code_color, bold_color):
    for i, part in enumerate(txt.split("\n")):
        if i:
            _line_break(p)
        if part == "":
            continue
        # the check mark is not in Segoe UI / Consolas -> use the symbol font
        for j, chunk in enumerate(part.split("\u2714")):
            if j:
                r = p.add_run()
                r.text = "\u2714"
                r.font.bold = bold
                r.font.name = "Segoe UI Symbol"
                r.font.size = Pt(size)
                r.font.italic = italic
                r.font.color.rgb = color
            if chunk == "":
                continue
            r = p.add_run()
            r.text = chunk
            r.font.bold = bold
            r.font.name = font
            r.font.size = Pt(size)
            r.font.italic = italic
            r.font.color.rgb = color


def rich(p, txt, size=15, color=INK, font=FONT, bold=False,
         code_color=None, bold_color=None, italic=False):
    """Adds runs to p, honouring **bold**, `code` and \\n (line break)."""
    for tok in INLINE_RE.split(txt):
        if not tok:
            continue
        if tok.startswith("**") and tok.endswith("**"):
            _emit(p, tok[2:-2], size, bold_color or color, font, True, italic,
                  code_color, bold_color)
        elif tok.startswith("`") and tok.endswith("`"):
            _emit(p, tok[1:-1], size * 0.94, code_color or TEAL, MONO, False,
                  False, code_color, bold_color)
        else:
            _emit(p, tok, size, color, font, bold, italic, code_color, bold_color)
    return p


def text(slide, x, y, w, h, content, size=15, color=INK, font=FONT, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, line_spacing=1.12,
         italic=False):
    tf = tbox(slide, x, y, w, h)
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    p.line_spacing = line_spacing
    rich(p, content, size=size, color=color, font=font, bold=bold, italic=italic)
    return tf


# --------------------------------------------------------------------------
# C++ syntax highlighter
# --------------------------------------------------------------------------
TOKEN_RE = re.compile(
    r'(?P<cmt>//[^\n]*)'
    r'|(?P<pre>\#\s*[A-Za-z_]+)'
    r'|(?P<str>"(?:[^"\\\n]|\\.)*")'
    r'|(?P<chr>\'(?:[^\'\\\n]|\\.)*\')'
    r'|(?P<word>[A-Za-z_][A-Za-z_0-9]*)'
    r'|(?P<num>\b\d+\b)'
    r'|(?P<sp>\s+)'
    r'|(?P<oth>.)'
)


def highlight(line):
    out = []
    for m in TOKEN_RE.finditer(line):
        kind, txt = m.lastgroup, m.group()
        if kind == "cmt":
            out.append((txt, C_CODE["cmt"]))
        elif kind == "pre":
            out.append((txt, C_CODE["pre"]))
        elif kind in ("str", "chr"):
            out.append((txt, C_CODE["str"]))
        elif kind == "word":
            if txt in KEYWORDS:
                col = C_CODE["kw"]
            elif txt[0].isupper():
                col = C_CODE["type"]
            else:
                col = C_CODE["def"]
            out.append((txt, col))
        elif kind == "num":
            out.append((txt, C_CODE["num"]))
        else:
            out.append((txt, C_CODE["def"]))
    return out


def code_panel(slide, x, y, w, code, max_h=None, size=11.5, min_size=6.5,
               title=None, term=False, pad=0.15):
    """Dark code / terminal panel. Returns (shape, height_in_inches)."""
    lines = code.split("\n")
    bg = TERM_BG if term else CODE_BG
    head_h = 0.32 if title else 0.0
    wpx = (w / 914400.0) * 96.0
    avail_px = wpx - pad * 2 * 96.0

    def needed(sz):
        total = 0
        for ln in lines:
            n = count_lines([(ln, sz, MONO, False)], avail_px) if ln.strip() else 1
            total += n * sz * LH_CODE / 72.0 * 96.0
        return (total / 96.0) + pad * 2 + head_h

    if max_h is not None:
        lim = max_h / 914400.0
        while size > min_size and needed(size) > lim:
            size -= 0.25

    h = needed(size)
    h_emu = Emu(int(h * 914400))

    panel = rect(slide, x, y, w, h_emu, fill=bg,
                 line=LINE if not term else None, lw=0.75,
                 shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.03)

    if title:
        rect(slide, x, y, w, Inches(head_h), fill=NAVY_3 if not term else NAVY_2)
        text(slide, x + Inches(0.14), y + Inches(0.045), w - Inches(0.28),
             Inches(0.24), title, size=9.5,
             color=RGBColor(0xA8, 0xBC, 0xD6), font=MONO, bold=True)

    tf = tbox(slide, x + Inches(pad), y + Inches(pad + head_h),
              w - Inches(pad * 2), Inches(h - pad * 2 - head_h))
    first = True
    for ln in lines:
        p = tf.paragraphs[0] if first else tf.add_paragraph()
        first = False
        p.line_spacing = 1.18
        no_bullet(p)
        if not ln.strip():
            r = p.add_run()
            r.text = " "
            r.font.size = Pt(size)
            r.font.name = MONO
            continue
        for txt, col in highlight(ln):
            r = p.add_run()
            r.text = txt
            r.font.name = MONO
            r.font.size = Pt(size)
            r.font.color.rgb = col
            if col == C_CODE["cmt"]:
                r.font.italic = True
    return panel, h


# --------------------------------------------------------------------------
# Page chrome
# --------------------------------------------------------------------------
PAGE_NO = [0]


def bg_page(slide, color=PAGE):
    rect(slide, 0, 0, SW, SH, fill=color)


def footer(slide, label=None):
    PAGE_NO[0] += 1
    if label is None:
        label = FOOTER_LABEL
    rect(slide, ML, SH - Inches(0.62), CW, Emu(9525), fill=LINE)
    text(slide, ML, SH - Inches(0.5), CW * 0.75, Inches(0.26), label,
         size=9.5, color=GREY_L, font=FONT)
    text(slide, SW - MR - Inches(1.2), SH - Inches(0.5), Inches(1.2), Inches(0.26),
         "%02d" % PAGE_NO[0], size=9.5, color=GREY_L, font=FONT,
         align=PP_ALIGN.RIGHT, bold=True)


def header(slide, title, kicker=None, accent=ORANGE):
    if kicker:
        text(slide, ML, Inches(0.44), CW, Inches(0.26), kicker.upper(),
             size=10.5, color=accent, font=FONT_B, bold=True)
        ty = Inches(0.70)
    else:
        ty = Inches(0.52)
    size = 27
    while size > 18 and text_px(title, size, FONT_B, True) > (CW / 914400.0) * 96.0:
        size -= 0.5
    text(slide, ML, ty, CW, Inches(0.5), title, size=size, color=NAVY,
         font=FONT_B, bold=True)
    rect(slide, ML, ty + Inches(0.48), Inches(0.82), Inches(0.055), fill=accent)


def content_slide(prs, title, kicker=None, accent=ORANGE):
    s = blank(prs)
    bg_page(s)
    header(s, title, kicker, accent)
    footer(s)
    return s


def blank(prs):
    return prs.slides.add_slide(prs.slide_layouts[6])


def section_slide(prs, number, title, subtitle, items=None):
    s = blank(prs)
    bg_page(s, NAVY)
    rect(s, SW - Inches(3.1), 0, Inches(3.1), SH, fill=NAVY_2)
    rect(s, SW - Inches(3.1), 0, Inches(0.045), SH, fill=ORANGE)
    for i in range(6):
        rect(s, SW - Inches(2.6) + Inches(0.34) * i, Inches(1.0),
             Inches(0.14), Inches(0.14), fill=NAVY_3, shape=MSO_SHAPE.OVAL)
    text(s, ML, Inches(1.65), Inches(4.0), Inches(1.6), number, size=104,
         color=NAVY_3, font=FONT_B, bold=True)
    rect(s, ML, Inches(3.30), Inches(0.9), Inches(0.06), fill=ORANGE)
    text(s, ML, Inches(3.55), Inches(8.2), Inches(0.8), title, size=38, color=WHITE,
         font=FONT_B, bold=True)
    text(s, ML, Inches(4.38), Inches(7.8), Inches(0.5), subtitle, size=14.5,
         color=RGBColor(0x9F, 0xB3, 0xCC))
    if items:
        y = Inches(5.05)
        for it in items:
            rect(s, ML, y + Inches(0.08), Inches(0.09), Inches(0.09), fill=ORANGE,
                 shape=MSO_SHAPE.OVAL)
            text(s, ML + Inches(0.3), y, Inches(8.0), Inches(0.3), it, size=12.5,
                 color=RGBColor(0xC3, 0xD1, 0xE2))
            y += Inches(0.36)
    PAGE_NO[0] += 1
    return s


# --------------------------------------------------------------------------
# Bullets
# --------------------------------------------------------------------------
def bullets(slide, x, y, w, items, size=14, gap=0.15, max_h=None,
            color=INK, bullet_color=None, min_size=10.5):
    bullet_color = bullet_color or ORANGE
    norm = [(it if isinstance(it, tuple) else (0, it)) for it in items]
    wpx = (w / 914400.0) * 96.0
    if max_h is None:
        max_h = FOOT_TOP - y - Inches(0.12)
    lim = max_h / 914400.0 * 96.0

    def measure(sz, g):
        h = 0.0
        for lvl, txt in norm:
            ind = (0.30 * lvl + 0.26) * 96.0
            n = count_lines(plain_runs(txt, sz, FONT), wpx - ind)
            h += n * sz * LH_TEXT / 72.0 * 96.0 + g * 96.0
        return h

    while size > min_size and measure(size, gap) > lim:
        size -= 0.5
        gap = max(0.07, gap - 0.012)

    yy = y
    for lvl, txt in norm:
        ind = Inches(0.30 * lvl)
        avail = wpx - (0.30 * lvl + 0.26) * 96.0
        n = count_lines(plain_runs(txt, size, FONT), avail)
        hpx = n * size * LH_TEXT / 72.0 * 96.0
        h = Emu(int(hpx * 914400 / 96))
        tf = tbox(slide, x + ind, yy, w - ind, h)
        p = tf.paragraphs[0]
        p.line_spacing = 1.12
        if lvl == 0:
            set_bullet(p, "\u25aa", "%02X%02X%02X" % (bullet_color[0], bullet_color[1],
                                                       bullet_color[2]), marL=0.26)
        else:
            set_bullet(p, "\u2013", "%02X%02X%02X" % (GREY[0], GREY[1], GREY[2]),
                       marL=0.24)
        rich(p, txt, size=size, color=color if lvl == 0 else GREY)
        yy += h + Inches(gap)
    return yy


def callout(slide, x, y, w, title, body, accent=ORANGE, bg=ORANGE_L, size=12,
            h=None):
    wpx = (w / 914400.0) * 96.0
    inner = wpx - 0.44 * 96.0
    th = count_lines(plain_runs(title, size, FONT_B, True), inner)
    bh = count_lines(plain_runs(body, size - 0.5, FONT), inner) if body else 0
    if h is None:
        h = Inches((th * size * LH_TEXT / 72.0) +
                   (bh * (size - 0.5) * LH_TEXT / 72.0 if body else 0) + 0.30)
    rect(slide, x, y, w, h, fill=bg, shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06)
    rect(slide, x, y, Inches(0.055), h, fill=accent)
    tf = tbox(slide, x + Inches(0.24), y + Inches(0.13), w - Inches(0.44),
              h - Inches(0.22))
    p = tf.paragraphs[0]
    p.line_spacing = 1.12
    rich(p, title, size=size, color=accent, font=FONT_B, bold=True)
    if body:
        p2 = tf.add_paragraph()
        p2.line_spacing = 1.12
        p2.space_before = Pt(3)
        rich(p2, body, size=size - 0.5, color=INK)
    return h


# --------------------------------------------------------------------------
# Table
# --------------------------------------------------------------------------
def table(slide, x, y, w, data, col_w=None, size=12.5, row_h=0.4,
          head_h=0.44, head_bg=NAVY, zebra=RGBColor(0xF0, 0xF3, 0xF8)):
    rows, cols = len(data), len(data[0])
    gf = slide.shapes.add_table(rows, cols, int(x), int(y), int(w),
                                int(Inches(head_h + row_h * (rows - 1))))
    tbl = gf.table

    tblel = gf._element.graphic.graphicData.tbl
    tblPr = tblel.find(qn("a:tblPr"))
    for el in tblPr.findall(qn("a:tableStyleId")):
        tblPr.remove(el)
    tblPr.append(parse_xml(
        '<a:tableStyleId %s>{2D5ABB26-0587-4C30-8999-92F81FD0307C}</a:tableStyleId>'
        % nsdecls("a")))
    tbl.first_row = False
    tbl.horz_banding = False

    if col_w:
        tot = float(sum(col_w))
        for i, cwv in enumerate(col_w):
            tbl.columns[i].width = int(w * cwv / tot)
    tbl.rows[0].height = Inches(head_h)
    for r in range(1, rows):
        tbl.rows[r].height = Inches(row_h)

    for r in range(rows):
        for c in range(cols):
            cell = tbl.cell(r, c)
            cell.margin_left = Inches(0.11)
            cell.margin_right = Inches(0.09)
            cell.margin_top = Inches(0.04)
            cell.margin_bottom = Inches(0.04)
            cell.vertical_anchor = MSO_ANCHOR.MIDDLE
            cell.fill.solid()
            cell.fill.fore_color.rgb = head_bg if r == 0 else (
                WHITE if r % 2 else zebra)
            tf = cell.text_frame
            tf.word_wrap = True
            p = tf.paragraphs[0]
            p.line_spacing = 1.10
            val = data[r][c]
            chunks = str(val).split("\n")
            first = True
            for chunk in chunks:
                if not first:
                    p = tf.add_paragraph()
                    p.line_spacing = 1.10
                first = False
                if r == 0:
                    rich(p, chunk, size=size, color=WHITE, font=FONT_B, bold=True)
                else:
                    rich(p, chunk, size=size, color=INK)
            tcPr = cell._tc.get_or_add_tcPr()
            for tag in ("a:lnL", "a:lnR", "a:lnT", "a:lnB"):
                for el in tcPr.findall(qn(tag)):
                    tcPr.remove(el)
            bc = "0E1B2E" if r == 0 else "D9DFEA"
            for i, tag in enumerate(("a:lnL", "a:lnR", "a:lnT", "a:lnB")):
                tcPr.insert(i, parse_xml(
                    '<%s %s w="9525" cap="flat" cmpd="sng" algn="ctr">'
                    '<a:solidFill><a:srgbClr val="%s"/></a:solidFill>'
                    '<a:prstDash val="solid"/></%s>'
                    % (tag, nsdecls("a"), bc, tag)))
    return gf


# --------------------------------------------------------------------------
# UML-ish class boxes / arrows
# --------------------------------------------------------------------------
def class_box(slide, x, y, w, h, name, body=None, head=NAVY_2, body_bg=WHITE,
              accent=ORANGE, size=12, name_size=14):
    """h is a length (EMU)."""
    rect(slide, x, y, w, h, fill=body_bg, line=LINE, lw=1.0,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.06)
    rect(slide, x, y, w, Inches(0.40), fill=head,
         shape=MSO_SHAPE.ROUNDED_RECTANGLE, adj=0.30)
    rect(slide, x, y + Inches(0.20), w, Inches(0.20), fill=head)
    rect(slide, x, y, Inches(0.05), h, fill=accent)
    text(slide, x + Inches(0.18), y + Inches(0.09), w - Inches(0.3), Inches(0.26),
         name, size=name_size, color=WHITE, font=FONT_B, bold=True)
    if body:
        tf = tbox(slide, x + Inches(0.18), y + Inches(0.50), w - Inches(0.3),
                  h - Inches(0.56))
        first = True
        for ln in body:
            p = tf.paragraphs[0] if first else tf.add_paragraph()
            first = False
            p.line_spacing = 1.14
            rich(p, ln, size=size, color=INK, font=MONO)
    return None


def inherit_arrow(slide, x, y_base, y_tip, color=NAVY_2, dashed=False, w=1.5):
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x), int(y_tip),
                                     int(x), int(y_base))
    ln.line.color.rgb = color
    ln.line.width = Pt(w)
    if dashed:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                 int(x - Inches(0.10)), int(y_tip),
                                 int(Inches(0.20)), int(Inches(0.18)))
    tri.fill.background()
    tri.line.color.rgb = color
    tri.line.width = Pt(w)
    if dashed:
        tri.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    no_shadow(tri)
    return ln


def arrow_to(slide, x1, y1, x2, y2, color=NAVY_2, dashed=False, w=1.5):
    """UML inheritance arrow from (x1,y1) to (x2,y2) with a hollow triangle
    at the target end."""
    import math
    ln = slide.shapes.add_connector(MSO_CONNECTOR.STRAIGHT, int(x1), int(y1),
                                     int(x2), int(y2))
    ln.line.color.rgb = color
    ln.line.width = Pt(w)
    if dashed:
        ln.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    dx, dy = float(x2 - x1), float(y2 - y1)
    L = math.hypot(dx, dy) or 1.0
    ux, uy = dx / L, dy / L
    th = Inches(0.20)
    tw = Inches(0.20)
    ccx = x2 - ux * th / 2
    ccy = y2 - uy * th / 2
    tri = slide.shapes.add_shape(MSO_SHAPE.ISOSCELES_TRIANGLE,
                                 int(ccx - tw / 2), int(ccy - th / 2),
                                 int(tw), int(th))
    tri.fill.background()
    tri.line.color.rgb = color
    tri.line.width = Pt(w)
    if dashed:
        tri.line.dash_style = MSO_LINE_DASH_STYLE.DASH
    tri.rotation = math.degrees(math.atan2(ux, -uy))
    no_shadow(tri)
    return ln
